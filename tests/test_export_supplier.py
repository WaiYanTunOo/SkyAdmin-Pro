"""Excel export — supplier sheets match database schema."""

import pytest

from skyadmin_pro.database import Database
from skyadmin_pro.services.export import export_to_excel


@pytest.fixture
def db(tmp_path):
    return Database(tmp_path / "export.db")


def test_export_supplier_columns(db, tmp_path):
    supplier_id = db.add_supplier(
        name="IRD Liaison",
        company_name="Gov Services",
        contact="officer@example.com",
        notes="Tax filings",
    )
    client_id = db.get_or_create_client("Beta Co")
    db.add_supplier_payment(
        supplier_id=supplier_id,
        client_id=client_id,
        amount="2000",
        due_date="2026-10-01",
        notes="Filing fee",
    )
    db.add_supplier_service(
        supplier_id=supplier_id,
        company_name="Beta Co",
        service_type="Registered address",
        expiry_date="2026-12-31",
    )

    dest = tmp_path / "export.xlsx"
    export_to_excel(db, dest)
    assert dest.exists()

    import openpyxl

    wb = openpyxl.load_workbook(dest, read_only=True, data_only=True)
    try:
        suppliers = wb["Suppliers"]
        headers = [cell.value for cell in next(suppliers.iter_rows(min_row=1, max_row=1))]
        first = next(suppliers.iter_rows(min_row=2, max_row=2, values_only=True))
    finally:
        wb.close()
    assert "Company" in headers
    assert "Contact" in headers
    assert first[headers.index("Supplier")] == "IRD Liaison"
    assert first[headers.index("Company")] == "Gov Services"
    assert first[headers.index("Contact")] == "officer@example.com"

    wb = openpyxl.load_workbook(dest, read_only=True, data_only=True)
    try:
        payments = wb["Supplier Payments"]
        pay_headers = [cell.value for cell in next(payments.iter_rows(min_row=1, max_row=1))]
        pay_first = next(payments.iter_rows(min_row=2, max_row=2, values_only=True))
    finally:
        wb.close()
    assert "Notes" in pay_headers
    assert pay_first[pay_headers.index("Notes")] == "Filing fee"
